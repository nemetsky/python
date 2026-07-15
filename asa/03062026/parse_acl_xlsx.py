#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import re
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import sys

print("=" * 80)
print("Cisco ASA ACL Parser v6.0 - FINAL")
print("=" * 80)

try:
    with open('ACL_EXT.txt', 'r', encoding='utf-8') as f:
        ext_lines = f.readlines()
    with open('ACL_BRIEF.txt', 'r', encoding='utf-8') as f:
        brief_lines = f.readlines()
    print(f"EXT: {len(ext_lines)} lines")
    print(f"BRIEF: {len(brief_lines)} lines")
except FileNotFoundError as e:
    print(f"Error: {e}")
    sys.exit(1)

has_hex = any(re.search(r'0x[0-9a-fA-F]+\s*$', line) for line in ext_lines)
has_last_hit = any('Last Hit=' in line for line in brief_lines)

if has_hex and not has_last_hit:
    format_type = 'NEW'
    print("\nFormat: NEW (hex hashes)")
else:
    format_type = 'OLD'
    print("\nFormat: OLD (Last Hit in BRIEF)")

def hex_to_datetime(hex_ts):
    try:
        hex_ts = hex_ts.lower().replace('0x', '')
        ts_int = int(hex_ts, 16)
        
        if 1000000000 <= ts_int < 2147483647:
            dt = datetime.utcfromtimestamp(ts_int)
            return dt.strftime('%b %d %Y'), dt.strftime('%H:%M:%S')
        return 'N/A', 'N/A'
    except:
        return 'Error', 'N/A'

brief_data = {}

print("\nParsing BRIEF file...")
for line in brief_lines:
    if 'elements' in line and 'name hash' in line:
        continue
    
    line_match = re.search(r'line\s+(\d+)', line)
    if not line_match:
        tokens = line.strip().split()
        if len(tokens) == 4:
            try:
                hash1 = tokens[0].lower()
                hash2 = tokens[1].lower()
                hitcount = int(tokens[2], 16)
                timestamp_hex = tokens[3].lower()
                
                date_str, time_str = hex_to_datetime(timestamp_hex)
                
                brief_data[hash1] = {'hitcount': hitcount, 'date': date_str, 'time': time_str}
                brief_data[hash2] = {'hitcount': hitcount, 'date': date_str, 'time': time_str}
            except:
                pass
        continue
    
    line_num = int(line_match.group(1))
    
    hitcnt = 0
    hitcnt_match = re.search(r'\(hitcnt=(\d+)\)', line)
    if hitcnt_match:
        hitcnt = int(hitcnt_match.group(1))
    
    date_str = ''
    time_str = ''
    lasthit_match = re.search(r'Last Hit=(\d+:\d+:\d+)\s+UTC\s+(\w+\s+\d+\s+\d+)', line)
    if lasthit_match:
        time_str = lasthit_match.group(1)
        date_str = lasthit_match.group(2)
    
    brief_data[line_num] = {'hitcount': hitcnt, 'date': date_str, 'time': time_str}

print(f"Records in brief_data: {len(brief_data)}")

def parse_address(tokens, idx):
    if idx >= len(tokens):
        return '', idx
    if tokens[idx] == 'any':
        return 'any', idx + 1
    elif tokens[idx] == 'host':
        if idx + 1 < len(tokens):
            return f"host {tokens[idx + 1]}", idx + 2
    elif tokens[idx] == 'object-group':
        if idx + 1 < len(tokens):
            return f"object-group {tokens[idx + 1]}", idx + 2
    elif tokens[idx] == 'object':
        if idx + 1 < len(tokens):
            return f"object {tokens[idx + 1]}", idx + 2
    else:
        addr = tokens[idx]
        if idx + 1 < len(tokens) and re.match(r'^\d+\.\d+\.\d+\.\d+$', tokens[idx + 1]):
            addr += f" {tokens[idx + 1]}"
            return addr, idx + 2
        return addr, idx + 1
    return tokens[idx], idx + 1

def parse_port(tokens, idx):
    if idx >= len(tokens):
        return '', idx
    if tokens[idx] in ['eq', 'lt', 'gt', 'neq']:
        if idx + 1 < len(tokens):
            return f"{tokens[idx]} {tokens[idx + 1]}", idx + 2
    elif tokens[idx] == 'range':
        if idx + 2 < len(tokens):
            return f"range {tokens[idx + 1]} {tokens[idx + 2]}", idx + 3
    return '', idx

rules = []

print("\nParsing EXT file...")
for line in ext_lines:
    if 'access-list' not in line or 'elements' in line:
        continue
    
    rule = {
        'ACL_Name': '', 'Line': 0, 'Action': '', 'Protocol': '',
        'Source': '', 'Source_Port': '', 'Destination': '', 'Destination_Port': '',
        'Hit_Count': 0, 'Date_Last_Hit': '', 'Time_Last_Hit': '',
        'Unique_Hash': '', 'Remark': ''
    }
    
    acl_match = re.search(r'access-list\s+(\S+)', line)
    if acl_match:
        rule['ACL_Name'] = acl_match.group(1)
    
    line_match = re.search(r'line\s+(\d+)', line)
    if line_match:
        rule['Line'] = int(line_match.group(1))
    else:
        continue
    
    if ' remark ' in line:
        remark_match = re.search(r'remark\s+(.+?)(?:\s*\(|\s*$)', line)
        if remark_match:
            rule['Remark'] = remark_match.group(1).strip()
        rules.append(rule)
        continue
    
    hitcnt_match = re.search(r'\(hitcnt=(\d+)\)', line)
    if hitcnt_match:
        rule['Hit_Count'] = int(hitcnt_match.group(1))
    
    hash_match = re.search(r'(0x[0-9a-fA-F]+)\s*$', line)
    if hash_match:
        rule['Unique_Hash'] = hash_match.group(1).lower()
    
    if ' extended ' in line:
        clean_line = re.sub(r'\s*\(hitcnt=\d+\).*', '', line)
        clean_line = re.sub(r'\s*\(Last Hit=.*\)', '', clean_line)
        clean_line = re.sub(r'\s+0x[0-9a-fA-F]+\s*$', '', clean_line)
        
        tokens = clean_line.split()
        try:
            ext_idx = tokens.index('extended')
            tokens = tokens[ext_idx + 1:]
        except:
            pass
        
        if len(tokens) >= 3:
            rule['Action'] = tokens[0]
            rule['Protocol'] = tokens[1]
            idx = 2
            rule['Source'], idx = parse_address(tokens, idx)
            rule['Source_Port'], idx = parse_port(tokens, idx)
            rule['Destination'], idx = parse_address(tokens, idx)
            rule['Destination_Port'], idx = parse_port(tokens, idx)
    
    if format_type == 'OLD':
        if rule['Line'] in brief_data:
            rule['Hit_Count'] = brief_data[rule['Line']].get('hitcount', rule['Hit_Count'])
            rule['Date_Last_Hit'] = brief_data[rule['Line']].get('date', '')
            rule['Time_Last_Hit'] = brief_data[rule['Line']].get('time', '')
    else:
        if rule['Unique_Hash']:
            hash_key = rule['Unique_Hash'].replace('0x', '')
            if hash_key in brief_data:
                rule['Hit_Count'] = brief_data[hash_key].get('hitcount', rule['Hit_Count'])
                rule['Date_Last_Hit'] = brief_data[hash_key].get('date', 'N/A')
                rule['Time_Last_Hit'] = brief_data[hash_key].get('time', 'N/A')
    
    rules.append(rule)

print(f"Rules found: {len(rules)}")

wb = Workbook()
ws = wb.active
ws.title = "ACL Statistics"

headers = ['ACL_Name', 'Line', 'Action', 'Protocol', 'Source', 'Source_Port',
           'Destination', 'Destination_Port', 'Hit_Count', 'Date_Last_Hit', 
           'Time_Last_Hit', 'Unique_Hash', 'Remark']
ws.append(headers)

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')

for rule in sorted(rules, key=lambda x: (x['ACL_Name'], x['Line'])):
    row = [rule['ACL_Name'], rule['Line'], rule['Action'], rule['Protocol'],
           rule['Source'], rule['Source_Port'], rule['Destination'], 
           rule['Destination_Port'], rule['Hit_Count'], rule['Date_Last_Hit'], 
           rule['Time_Last_Hit'], rule['Unique_Hash'], rule['Remark']]
    ws.append(row)
    
    hit_cell = ws.cell(row=ws.max_row, column=9)
    if rule['Hit_Count'] == 0 and not rule['Remark']:
        hit_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        hit_cell.font = Font(color="9C0006", bold=True)
    elif rule['Hit_Count'] > 10000:
        hit_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        hit_cell.font = Font(color="006100", bold=True)
    elif rule['Hit_Count'] > 100:
        hit_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

for col, width in {'A':15,'B':6,'C':8,'D':10,'E':35,'F':12,'G':35,'H':15,
                   'I':12,'J':15,'K':15,'L':12,'M':30}.items():
    ws.column_dimensions[col].width = width

wb.save('ACL_Report.xlsx')
print("\nFile saved: ACL_Report.xlsx")

print("\n" + "=" * 80)
zero_hit = [r for r in rules if r['Hit_Count'] == 0 and not r['Remark']]
with_dates = [r for r in rules if r['Date_Last_Hit'] and r['Date_Last_Hit'] != 'N/A' and not r['Remark']]

print(f"Total rules: {len([r for r in rules if not r['Remark']])}")
print(f"Rules with dates: {len(with_dates)}")
print(f"Inactive (=0): {len(zero_hit)}")

if with_dates:
    print("\nLast Hit date examples:")
    for r in [x for x in with_dates][:3]:
        print(f"  Line {r['Line']}: {r['Date_Last_Hit']} {r['Time_Last_Hit']} (Hit: {r['Hit_Count']})")

if zero_hit:
    print("\nCandidates for removal:")
    for r in zero_hit[:5]:
        print(f"  - {r['ACL_Name']} line {r['Line']}: {r['Hit_Count']} hits")

print("\nDONE!")
print("=" * 80)