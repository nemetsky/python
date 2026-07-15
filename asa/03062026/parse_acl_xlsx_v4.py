#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import re
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone
import sys

print("=" * 80)
print("Cisco ASA ACL Parser v8.0 - No Remarks")
print("=" * 80)

EXT_FILE = 'ACL_EXT.txt'
BRIEF_FILE = 'ACL_BRIEF.txt'
OUTPUT_FILE = 'ACL_Report.xlsx'

try:
    with open(EXT_FILE, 'r', encoding='utf-8') as f:
        ext_lines = f.readlines()
    print(f"EXT: {len(ext_lines)} lines")
    try:
        with open(BRIEF_FILE, 'r', encoding='utf-8') as f:
            brief_lines = f.readlines()
        print(f"BRIEF: {len(brief_lines)} lines")
    except FileNotFoundError:
        print("BRIEF: File not found (will use hitcnt from EXT)")
        brief_lines = []
except FileNotFoundError as e:
    print(f"Error: {e}")
    sys.exit(1)

has_hex = any(re.search(r'0x[0-9a-fA-F]+\s*$', line) for line in ext_lines)
has_last_hit = any('Last Hit=' in line for line in brief_lines)
format_type = 'NEW' if has_hex and not has_last_hit else 'OLD'
print(f"\nFormat: {format_type} ({'hex hashes' if format_type == 'NEW' else 'Last Hit in BRIEF'})")

regex_patterns = {
    'line_num': re.compile(r'line\s+(\d+)'),
    'acl_name': re.compile(r'access-list\s+(\S+)'),
    'hitcnt': re.compile(r'\(hitcnt=(\d+)\)'),
    'hash': re.compile(r'(0x[0-9a-fA-F]+)\s*$'),
    'lasthit': re.compile(r'Last Hit=(\d+:\d+:\d+)\s+UTC\s+(\w+\s+\d+\s+\d+)'),
    'remark': re.compile(r'remark\s+(.+?)(?:\s*\(|\s*$)'),
    'ipv4': re.compile(r'^\d+\.\d+\.\d+\.\d+$'),
}

def hex_to_datetime(hex_ts):
    try:
        hex_ts = hex_ts.lower().replace('0x', '')
        ts_int = int(hex_ts, 16)
        if 1000000000 <= ts_int < 2147483647:
            dt = datetime.fromtimestamp(ts_int, tz=timezone.utc)
            return dt.strftime('%b %d %Y'), dt.strftime('%H:%M:%S')
        return 'N/A', 'N/A'
    except:
        return 'Error', 'N/A'

brief_data = {}
brief_hash_count = 0
brief_line_count = 0

print("\nParsing BRIEF file...")
for line in brief_lines:
    if not line.strip():
        continue
    if 'elements' in line and 'name hash' in line:
        continue
    line_match = regex_patterns['line_num'].search(line)
    if not line_match:
        tokens = line.strip().split()
        if len(tokens) == 4:
            try:
                hash1 = tokens[0].lower()
                hash2 = tokens[1].lower()
                hitcount = int(tokens[2], 16)
                timestamp_hex = tokens[3].lower()
                date_str, time_str = hex_to_datetime(timestamp_hex)
                brief_data[hash1] = {'hitcount': hitcount, 'date': date_str, 'time': time_str}
                brief_data[hash2] = {'hitcount': hitcount, 'date': date_str, 'time': time_str}
                brief_hash_count += 1
            except:
                pass
            continue
    line_num = int(line_match.group(1))
    hitcnt = 0
    hitcnt_match = regex_patterns['hitcnt'].search(line)
    if hitcnt_match:
        hitcnt = int(hitcnt_match.group(1))
    date_str = ''
    time_str = ''
    lasthit_match = regex_patterns['lasthit'].search(line)
    if lasthit_match:
        time_str = lasthit_match.group(1)
        date_str = lasthit_match.group(2)
    brief_data[str(line_num)] = {'hitcount': hitcnt, 'date': date_str, 'time': time_str}
    brief_line_count += 1

print(f"BRIEF hash pairs: {brief_hash_count}")
print(f"BRIEF line entries: {brief_line_count}")
print(f"Total brief_data records: {len(brief_data)}")

def parse_address(tokens, start_idx):
    if start_idx >= len(tokens):
        return '', start_idx
    if tokens[start_idx] == 'any':
        return 'any', start_idx + 1
    elif tokens[start_idx] == 'host':
        if start_idx + 1 < len(tokens):
            return f"host {tokens[start_idx + 1]}", start_idx + 2
    elif tokens[start_idx] == 'object-group':
        if start_idx + 1 < len(tokens):
            return f"object-group {tokens[start_idx + 1]}", start_idx + 2
    elif tokens[start_idx] == 'object':
        if start_idx + 1 < len(tokens):
            return f"object {tokens[start_idx + 1]}", start_idx + 2
    else:
        addr = tokens[start_idx]
        if start_idx + 1 < len(tokens) and regex_patterns['ipv4'].match(tokens[start_idx + 1]):
            addr += f" {tokens[start_idx + 1]}"
            return addr, start_idx + 2
        return addr, start_idx + 1
    return tokens[start_idx], start_idx + 1

def parse_port(tokens, start_idx):
    if start_idx >= len(tokens):
        return '', start_idx
    
    if tokens[start_idx] in ['eq', 'lt', 'gt', 'neq']:
        if start_idx + 1 < len(tokens):
            return f"{tokens[start_idx]} {tokens[start_idx + 1]}", start_idx + 2
    elif tokens[start_idx] == 'range':
        if start_idx + 2 < len(tokens):
            return f"range {tokens[start_idx + 1]} {tokens[start_idx + 2]}", start_idx + 3
    elif tokens[start_idx] == 'object-group':
        next_token_idx = start_idx + 2
        if next_token_idx < len(tokens):
            next_token = tokens[next_token_idx]
            if next_token in ['eq', 'range', 'lt', 'gt', 'neq']:
                 return '', start_idx
            else:
                 return f"object-group {tokens[start_idx + 1]}", start_idx + 2
        else:
             return '', start_idx

    return '', start_idx

rules_dict = {}
ext_hitcnt_count = 0
brief_match_count = 0
brief_no_match_count = 0
remarks_count = 0
current_remark = ""
current_acl_name = ""

print("\nParsing EXT file...")
for i, line in enumerate(ext_lines):
    if not line.strip():
        continue
    if 'access-list' not in line or 'elements' in line:
        continue
    if i % 10000 == 0:
        print(f"  Processed: {i}/{len(ext_lines)} lines...")
    
    acl_match = regex_patterns['acl_name'].search(line)
    if acl_match:
        new_acl = acl_match.group(1)
        if new_acl != current_acl_name:
            current_acl_name = new_acl
            current_remark = ""
            
    if ' remark ' in line:
        remark_text = line.split('remark', 1)[-1].strip()
        remark_text = re.sub(r'===[\|\s]*|[\|\s]*===', '', remark_text).strip()
        current_remark = remark_text
        remarks_count += 1
        continue

    clean_line = re.sub(r'\s*\(hitcnt=\d+\).*', '', line)
    clean_line = re.sub(r'\s*line\s+\d+', '', clean_line)
    clean_line = re.sub(r'\s+', ' ', clean_line)
    final_line = clean_line.strip()

    rule = {
        'ACL_Name': '', 'Line': 0, 'Action': '', 'Protocol': '', 'Object_Port': '',
        'Source': '', 'Destination': '', 'Destination_Port': '',
        'Hit_Count': 0, 'Date_Last_Hit': '', 'Time_Last_Hit': '',
        'Unique_Hash': '', 'Remark': current_remark,
        'Original_Line': final_line
    }
    
    if acl_match:
        rule['ACL_Name'] = acl_match.group(1)
    
    line_match = regex_patterns['line_num'].search(line)
    if line_match:
        rule['Line'] = int(line_match.group(1))
    else:
        continue
        
    hitcnt_match = regex_patterns['hitcnt'].search(line)
    if hitcnt_match:
        rule['Hit_Count'] = int(hitcnt_match.group(1))
        ext_hitcnt_count += 1
        
    hash_match = regex_patterns['hash'].search(line)
    if hash_match:
        rule['Unique_Hash'] = hash_match.group(1).lower()
        
    if ' extended ' in line or ' permit ' in line or ' deny ' in line:
        clean_line = re.sub(r'\s*\(hitcnt=\d+\).*', '', line)
        clean_line = re.sub(r'\s*\(Last Hit=.*\)', '', clean_line)
        clean_line = re.sub(r'\s+0x[0-9a-fA-F]+\s*$', '', clean_line)
        tokens = clean_line.split()
        try:
            if 'extended' in tokens:
                ext_idx = tokens.index('extended')
                tokens = tokens[ext_idx + 1:]
            elif 'permit' in tokens:
                ext_idx = tokens.index('permit')
                tokens = tokens[ext_idx:]
            elif 'deny' in tokens:
                ext_idx = tokens.index('deny')
                tokens = tokens[ext_idx:]
        except:
            pass
            
        if len(tokens) >= 3:
            rule['Action'] = tokens[0]
            
            if tokens[1] == 'object-group':
                rule['Protocol'] = ''
                rule['Object_Port'] = f"object-group {tokens[2]}"
                idx = 3
            else:
                rule['Protocol'] = tokens[1]
                idx = 2
            
            rule['Source'], idx = parse_address(tokens, idx)
            _, idx = parse_port(tokens, idx)
            rule['Destination'], idx = parse_address(tokens, idx)
            rule['Destination_Port'], idx = parse_port(tokens, idx)

    if brief_lines:
        if format_type == 'OLD':
            if str(rule['Line']) in brief_data:
                rule['Hit_Count'] = brief_data[str(rule['Line'])].get('hitcount', rule['Hit_Count'])
                rule['Date_Last_Hit'] = brief_data[str(rule['Line'])].get('date', '')
                rule['Time_Last_Hit'] = brief_data[str(rule['Line'])].get('time', '')
                brief_match_count += 1
            else:
                brief_no_match_count += 1
        else:
            if rule['Unique_Hash']:
                hash_key = rule['Unique_Hash'].replace('0x', '')
                if hash_key in brief_data:
                    rule['Hit_Count'] = brief_data[hash_key].get('hitcount', rule['Hit_Count'])
                    rule['Date_Last_Hit'] = brief_data[hash_key].get('date', 'N/A')
                    rule['Time_Last_Hit'] = brief_data[hash_key].get('time', 'N/A')
                    brief_match_count += 1
                else:
                    brief_no_match_count += 1
                    
    key = f"{rule['ACL_Name']}_{rule['Line']}"
    if key not in rules_dict:
        rules_dict[key] = rule

rules = list(rules_dict.values())
print(f"Rules found (unique lines): {len(rules)}")
print(f"Remarks skipped: {remarks_count}")
print(f"EXT hitcnt values: {ext_hitcnt_count}")
if brief_lines:
    print(f"BRIEF matches: {brief_match_count}")
    print(f"BRIEF no match: {brief_no_match_count}")

print("\nCreating Excel file...")
wb = Workbook()
ws = wb.active
ws.title = "ACL Statistics"

headers = ['ACL_Name', 'Line', 'Action', 'Protocol', 'Object_Port', 'Source', 'Destination', 'Destination_Port',
           'Hit_Count', 'Date_Last_Hit', 'Time_Last_Hit', 'Unique_Hash', 'Remark', 'Original_Line']
ws.append(headers)

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')

for rule in sorted(rules, key=lambda x: (x['ACL_Name'], x['Line'])):
    row = [rule['ACL_Name'], rule['Line'], rule['Action'], rule['Protocol'], rule['Object_Port'],
           rule['Source'], rule['Destination'], rule['Destination_Port'], rule['Hit_Count'], rule['Date_Last_Hit'],
           rule['Time_Last_Hit'], rule['Unique_Hash'], rule['Remark'], rule['Original_Line']]
    ws.append(row)
    hit_cell = ws.cell(row=ws.max_row, column=9)
    if rule['Hit_Count'] == 0:
        hit_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        hit_cell.font = Font(color="9C0006", bold=True)
    elif rule['Hit_Count'] > 10000:
        hit_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        hit_cell.font = Font(color="006100", bold=True)
    elif rule['Hit_Count'] > 100:
        hit_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

for col, width in {'A':15, 'B':6, 'C':8, 'D':10, 'E':25, 'F':35, 'G':35, 'H':15, 'I':12, 'J':15, 'K':15, 'L':12, 'M':40, 'N':120}.items():
    ws.column_dimensions[col].width = width

wb.save(OUTPUT_FILE)
print(f"File saved: {OUTPUT_FILE}")

print("\n" + "=" * 80)
zero_hit = [r for r in rules if r['Hit_Count'] == 0]
with_dates = [r for r in rules if r['Date_Last_Hit'] and r['Date_Last_Hit'] != 'N/A']
very_active = [r for r in rules if r['Hit_Count'] > 10000]
print(f"Total rules: {len(rules)}")
print(f"Rules with dates: {len(with_dates)}")
print(f"Inactive (=0): {len(zero_hit)}")
print(f"Very active (>10000): {len(very_active)}")
print(f"Remarks skipped: {remarks_count}")
if very_active:
    print("\nTOP-5 most active:")
    for i, rule in enumerate(sorted(very_active, key=lambda x: x['Hit_Count'], reverse=True)[:5], 1):
        print(f"  {i}. Line {rule['Line']}: {rule['Hit_Count']:,} hits")
if zero_hit:
    print(f"\nCandidates for removal ({len(zero_hit)} rules):")
    for r in sorted(zero_hit, key=lambda x: (x['ACL_Name'], x['Line']))[:10]:
        print(f"  - {r['ACL_Name']} line {r['Line']}: {r['Action']} {r['Protocol']}")
print("\n" + "=" * 80)
print("DONE!")
print("=" * 80)