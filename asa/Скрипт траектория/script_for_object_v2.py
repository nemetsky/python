#!/usr/bin/env python3
import re
import sys
import os
import argparse

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Отсутствует библиотека openpyxl. Установите: pip install openpyxl")
    sys.exit(1)

def parse_config(config_path):
    with open(config_path, 'r', encoding='utf-8', errors='ignore') as f:
        raw_lines = f.readlines()
    active_lines = [ln.rstrip('\r\n') for ln in raw_lines if not ln.lstrip().startswith('!') and ln.strip()]
    full_text = '\n'.join(active_lines)
    def_re = re.compile(r'^\s*(object|object-group)\s+(network|service|protocol|icmp-type)\s+(\S+)', re.IGNORECASE)
    defined = {}
    for line in active_lines:
        m = def_re.match(line)
        if m:
            obj_type, subtype, name = m.group(1), m.group(2), m.group(3)
            if name not in defined:
                defined[name] = {'type': obj_type, 'subtype': subtype, 'count': 0}
    for name in defined:
        escaped_name = re.escape(name)
        pattern = re.compile(r'(?<![A-Za-z0-9_.-])' + escaped_name + r'(?![A-Za-z0-9_.-])')
        defined[name]['count'] = len(pattern.findall(full_text))
    unused = [
        {'name': name, 'type': info['type'], 'subtype': info['subtype'], 'count': info['count']}
        for name, info in defined.items() if info['count'] == 1
    ]
    return sorted(unused, key=lambda x: x['name'])

def save_to_excel(unused_data, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Неиспользуемые объекты ASA"
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    cell_font = Font(size=11)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    headers = ["Имя объекта", "Тип", "Подтип", "Упоминаний", "Статус", "Рекомендация"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
    if unused_data:
        for row_idx, obj in enumerate(unused_data, 2):
            ws.cell(row=row_idx, column=1, value=obj['name']).font = cell_font
            ws.cell(row=row_idx, column=2, value=obj['type'].capitalize()).font = cell_font
            ws.cell(row=row_idx, column=3, value=obj['subtype'].capitalize()).font = cell_font
            ws.cell(row=row_idx, column=4, value=obj['count']).font = cell_font
            ws.cell(row=row_idx, column=5, value="Не используется в политиках/NAT/маршрутах").font = cell_font
            ws.cell(row=row_idx, column=6, value="Проверить контекст -> закомментировать на 7-14 дней -> удалить").font = cell_font
            for col in range(1, 7):
                ws.cell(row=row_idx, column=col).alignment = align_left
    else:
        ws.cell(row=2, column=1, value="Неиспользуемых объектов не найдено. Конфигурация чистая.").font = cell_font
    col_widths = [35, 15, 15, 12, 40, 55]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    wb.save(output_path)
    print(f"Результат сохранен в: {os.path.abspath(output_path)}")

def main():
    parser = argparse.ArgumentParser(
        description="Поиск неиспользуемых объектов Cisco ASA с экспортом в Excel",
        formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument("config", help="Путь к файлу running-config")
    parser.add_argument("-o", "--output", default="asa_unused_objects.xlsx",
                        help="Путь к выходному Excel-файлу (по умолчанию: asa_unused_objects.xlsx)")
    args = parser.parse_args()
    if not os.path.isfile(args.config):
        print(f"Файл конфигурации '{args.config}' не найден.")
        sys.exit(1)
    print(f"🔍 Анализ файла: {args.config}...")
    unused = parse_config(args.config)
    if not unused:
        print("Все объекты используются. Excel-файл создан с пустым списком.")
    else:
        print(f"Найдено {len(unused)} неиспользуемых объектов/групп.")
    save_to_excel(unused, args.output)
    print("Рекомендация: Перед удалением закомментируйте объект и понаблюдайте за трафиком 7–14 дней.")

if __name__ == '__main__':
    main()